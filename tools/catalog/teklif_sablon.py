# Price-offer input template (Turkish labels) for MODUNERA. Yellow cells = to be filled by the owner.
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

OR, DK, YL, LG = "BC3A14", "2E3B41", "FFF4C2", "F3E6DE"
thin = Side(style="thin", color="D9D0C4")
B = Border(left=thin, right=thin, top=thin, bottom=thin)
wb = Workbook()

def sheet(ws, title, note, headers, widths):
    ws.title = title
    ws["A1"] = title; ws["A1"].font = Font(bold=True, size=16, color=OR)
    ws["A2"] = note; ws["A2"].font = Font(italic=True, size=10, color="6E777B")
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=4, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=DK)
        c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = B
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[4].height = 32
    ws.freeze_panes = "A5"

def row(ws, r, vals, fill_cols=(), eur_cols=()):
    for i, v in enumerate(vals, 1):
        c = ws.cell(row=r, column=i, value=v); c.border = B; c.alignment = Alignment(wrap_text=True, vertical="top")
        if i in fill_cols: c.fill = PatternFill("solid", fgColor=YL)
        if i in eur_cols: c.number_format = '#,##0 "€"'

# 1 Fiyatlar
ws = wb.active
sheet(ws, "1 Fiyatlar", "Sarı hücrelere ANAHTAR TESLİM fiyatı (EUR, KDV hariç, fabrika çıkışı) yazın. Nakliye 4. sayfada ayrı. Yapılmayan kombinasyonu boş bırakın.",
      ["Model", "Uzunluk", "Kişi / yerleşim", "Standard €", "Standard Plus €", "Premium €", "Not"], [10, 10, 26, 15, 17, 15, 34])
models = [("MD 1", "1 Loft"), ("MD 2", "2 Loft"), ("MD 3", "Loft + oda"), ("MD 4", "Loft, oda, veranda"), ("MD 5", "Loft + opsiyonel oda"),
          ("MD 6", "Chalet"), ("MD 7", "1 Loft"), ("MD 8", "1 Loft"), ("MD 9", "1 Loft")]
lengths = ["6 m", "7 m", "8 m", "9 m", "10 m"]
def persons(L):
    return {"6 m": ["1 kişi"], "7 m": ["1 kişi", "2 kişi"], "8 m": ["1 kişi", "2–3 kişi"], "9 m": ["1 kişi", "3–4 kişi"], "10 m": ["1 kişi", "4–5 kişi"]}[L]
r = 5
for m, lay in models:
    for L in lengths:
        for p in persons(L):
            row(ws, r, [m, L, f"{p} · {lay}", None, None, None, ""], fill_cols=(4, 5, 6), eur_cols=(4, 5, 6)); r += 1
    r += 1
for sp in ["Spiegelhaus 7 m", "Spiegelhaus 8 m", "Spiegelhaus 9 m", "Spiegelhaus 10 m", "Bungalow (m² fiyatı)"]:
    row(ws, r, [sp.split(" ")[0], " ".join(sp.split(" ")[1:]), "", None, None, None, ""], fill_cols=(4, 5, 6), eur_cols=(4, 5, 6)); r += 1

# 2 Paket içerikleri
ws2 = wb.create_sheet()
sheet(ws2, "2 Paket İçerikleri", "Katalogdaki içerikler önceden dolduruldu. Yanlış/eksik olanı düzeltin, marka/model değişikliklerini yazın. Sarı = kontrol edin.",
      ["Kalem", "Standard", "Standard Plus", "Premium"], [24, 40, 40, 40])
items = [
 ("Şasi", "Profil çelik, O2 römork", "Galvaniz çelik veya profil", "Galvaniz çelik veya profil"),
 ("Yalıtım", "10 cm taş yünü", "10 cm taş yünü (Izocam)", "10 cm taş yünü (Izocam)"),
 ("Nem bariyeri", "–", "Onduline", "Onduline"),
 ("Çatı", "Onduline / kenet sac", "Onduline / kenet sac", "Onduline / kenet sac"),
 ("Cephe", "Thermowood / kenet sac", "Thermowood / kenet / kompozit", "Thermowood / kenet / kompozit"),
 ("İç duvar", "Kontrplak + duvar kağıdı", "Kontrplak, duvar kağıdı / AGT panel", "Marin kontrplak, duvar kağıdı / AGT panel"),
 ("Zemin", "8 mm laminat", "8 mm AGT laminat", "8 / 12 mm AGT laminat"),
 ("Pencere / kapı", "PVC", "PVC", "PVC / alüminyum"),
 ("Elektrik", "Anahtar, priz, LED", "Halojensiz, Viko", "Halojensiz, Viko, Philips/Samsung LED"),
 ("Vitrifiye / armatür", "Batarya, duş, asma klozet", "Daha üst segment", "Grohe batarya, Vitra seramik"),
 ("Mutfak", "Tezgah, eviye", "Masif tezgah, lake kapak", "Masif tezgah, lake / membran kapak"),
 ("Sabit mobilya", "MDF-Lam", "MDF-Lam, lake kapak", "MDF-Lam, lake / membran kapak"),
 ("Beyaz eşya", "–", "Beko: buzdolabı, çamaşır, bulaşık, TV, ocak, davlumbaz, termosifon", "Bosch: buzdolabı, çamaşır, bulaşık, TV, ocak, davlumbaz, termosifon"),
 ("Klima", "Altyapı", "Mitsubishi Electric / York", "Mitsubishi Electric / York"),
 ("Hareketli mobilya (yatak, kanepe, masa)", "?", "?", "?"),
 ("Yerden ısıtma", "?", "?", "?"),
 ("Güneş paneli", "?", "?", "?"),
 ("Akıllı ev", "?", "?", "?"),
 ("Ekstra", "–", "–", "Dekoratif pelet soba"),
]
for i, it in enumerate(items, 5):
    row(ws2, i, list(it), fill_cols=(2, 3, 4) if "?" in it else ())

# 3 Ekstralar
ws3 = wb.create_sheet()
sheet(ws3, "3 Ekstralar", "Opsiyonel ekstraların birim fiyatı (EUR, KDV hariç). Bir pakette zaten dahilse 'dahil' yazın.",
      ["Ekstra", "Açıklama", "Fiyat €", "Hangi pakette dahil?"], [28, 44, 14, 24])
ex = [("Yerden ısıtma", "Elektrikli / sulu, tüm zemin"), ("Güneş paneli + batarya", "kWp ve batarya kapasitesini yazın"), ("Akıllı ev", "Işık, ısıtma, klima uygulamadan"),
      ("Temiz su tankı", "Litre"), ("Gri su tankı / arıtma", "Biyolojik arıtma sistemi"), ("Uydu TV / internet (Setcom)", ""), ("Hareketli mobilya paketi", "Yatak, kanepe, masa, sandalye"),
      ("Veranda / teras", "m² fiyatı"), ("Cam tavan (MD 1)", ""), ("Pelet soba", ""), ("Ek uzunluk (+1 m)", "Model fark etmeksizin"), ("Tekerleksiz / temel üzeri", "Fark (+/-)")]
for i, (a, b) in enumerate(ex, 5):
    row(ws3, i, [a, b, None, ""], fill_cols=(3, 4), eur_cols=(3,))

# 4 Nakliye & montaj
ws4 = wb.create_sheet()
sheet(ws4, "4 Nakliye & Montaj", "Ev başına fiyatlar (EUR). Teslim noktası: örnek Büsum (DE). Başka ülke/şehir için satır ekleyin.",
      ["Kalem", "Açıklama", "Fiyat €", "Not"], [28, 44, 14, 30])
tr = [("Nakliye Sakarya → Büsum", "Ev başına, kendi aksı veya lowbed"), ("Gümrük / evrak (A.TR vb.)", "Ev başına"), ("Vinç / yerleştirme", "Sahada"),
      ("Montaj ekibi (opsiyonel)", "Kurulum, bağlantı hazırlığı"), ("Sigorta", "Nakliye sigortası")]
for i, (a, b) in enumerate(tr, 5):
    row(ws4, i, [a, b, None, ""], fill_cols=(3,), eur_cols=(3,))

# 5 Koşullar
ws5 = wb.create_sheet()
sheet(ws5, "5 Koşullar", "Teklifte yazacak ticari koşullar. Sarı hücreleri doldurun / düzeltin.", ["Konu", "Değer"], [32, 60])
cond = [("Fiyatlar", "EUR, KDV hariç, fabrika çıkışı (EXW Sakarya) – nakliye ayrı"), ("Peşinat", "%60 sipariş ile"), ("Kalan ödeme", "%40 teslimattan önce / teslimatta?"),
        ("Teklif geçerlilik süresi", "30 gün"), ("Üretim süresi", "? hafta"), ("Garanti", "? yıl / kapsam"), ("Toplu alım indirimi", "örn. 10+ ev için %?"), ("Müşteri", "örn. Büsum projesi – yatırımcı adı")]
for i, (a, b) in enumerate(cond, 5):
    row(ws5, i, [a, b], fill_cols=(2,))

out = str(__import__("pathlib").Path(__file__).parent / "MODUNERA-Teklif-Fiyat-Sablonu.xlsx")
import os; os.makedirs(os.path.dirname(out), exist_ok=True)
wb.save(out); print(out)
